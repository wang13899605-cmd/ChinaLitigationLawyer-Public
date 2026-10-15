#!/usr/bin/env python3
"""Build sanitized Word mothers for legal research and law compilations.

The input references are read-only.  Each output reuses the reference page
system, paragraph formatting, table geometry, and run styling while replacing
all matter-specific text with generic drafting slots.
"""

from __future__ import annotations

import argparse
import sys
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
UTILITIES = ROOT / "09_Tools" / "utilities"
sys.path.insert(0, str(UTILITIES))
from docx_template_sanitize import sanitize  # noqa: E402


W_T = qn("w:t")
W_R = qn("w:r")
W_RPR = qn("w:rPr")
W_PPR = qn("w:pPr")
W_SECTPR = qn("w:sectPr")
W_PAGE_BREAK_BEFORE = qn("w:pageBreakBefore")


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != W_SECTPR:
            body.remove(child)


def clean_paragraph_clone(source_paragraph, text: str, *, page_break_before: bool = False):
    source = source_paragraph._p
    paragraph = OxmlElement("w:p")
    if source.pPr is not None:
        ppr = deepcopy(source.pPr)
        if page_break_before and ppr.find(W_PAGE_BREAK_BEFORE) is None:
            ppr.append(OxmlElement("w:pageBreakBefore"))
        paragraph.append(ppr)
    elif page_break_before:
        ppr = OxmlElement("w:pPr")
        ppr.append(OxmlElement("w:pageBreakBefore"))
        paragraph.append(ppr)

    source_run = source.find(f".//{W_R}")
    run = OxmlElement("w:r")
    if source_run is not None:
        rpr = source_run.find(W_RPR)
        if rpr is not None:
            run.append(deepcopy(rpr))
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    paragraph.append(run)
    return paragraph


def clean_multirun_clone(source_paragraph, values: list[str]):
    source = source_paragraph._p
    paragraph = OxmlElement("w:p")
    if source.pPr is not None:
        paragraph.append(deepcopy(source.pPr))
    source_runs = source.findall(f".//{W_R}")
    if not source_runs:
        source_runs = [None]
    for index, value in enumerate(values):
        src = source_runs[min(index, len(source_runs) - 1)]
        run = OxmlElement("w:r")
        if src is not None:
            rpr = src.find(W_RPR)
            if rpr is not None:
                run.append(deepcopy(rpr))
        node = OxmlElement("w:t")
        node.text = value
        run.append(node)
        paragraph.append(run)
    return paragraph


def append_paragraph(document: Document, source_paragraph, text: str, *, page_break_before=False):
    body = document._element.body
    body.insert(
        len(body) - 1,
        clean_paragraph_clone(source_paragraph, text, page_break_before=page_break_before),
    )


def append_multirun(document: Document, source_paragraph, values: list[str]):
    body = document._element.body
    body.insert(len(body) - 1, clean_multirun_clone(source_paragraph, values))


def replace_cell_text(cell, text: str) -> None:
    paragraphs = cell._tc.findall(f".//{qn('w:p')}")
    if not paragraphs:
        return
    first = paragraphs[0]
    source_run = first.find(f".//{W_R}")
    ppr = deepcopy(first.find(W_PPR)) if first.find(W_PPR) is not None else None
    for child in list(first):
        first.remove(child)
    if ppr is not None:
        first.append(ppr)
    run = OxmlElement("w:r")
    if source_run is not None and source_run.find(W_RPR) is not None:
        run.append(deepcopy(source_run.find(W_RPR)))
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    first.append(run)
    for extra in paragraphs[1:]:
        parent = extra.getparent()
        if parent is not None:
            parent.remove(extra)


def append_case_table(document: Document, source_table, headers: list[str], rows: list[list[str]]) -> None:
    table_xml = deepcopy(source_table._tbl)
    document._element.body.insert(len(document._element.body) - 1, table_xml)
    table = document.tables[-1]
    while len(table.rows) > 2:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(rows) + 1:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    for idx, value in enumerate(headers):
        replace_cell_text(table.rows[0].cells[idx], value)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    if header_properties.find(qn("w:tblHeader")) is None:
        header_properties.append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            replace_cell_text(table.rows[row_index].cells[column_index], value)


def save_sanitized(document: Document, output: Path, title: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=output.parent) as handle:
        temporary = Path(handle.name)
    try:
        document.save(temporary)
        sanitize(
            temporary,
            output,
            {
                "title": title,
                "parts": {},
                "remove_unlisted_non_table": False,
                "blank_headers": True,
                "blank_notes": True,
            },
        )
    finally:
        temporary.unlink(missing_ok=True)


def build_four_column(source: Path, output: Path) -> None:
    document = Document(source)
    source_paragraphs = list(document.paragraphs)
    source_table = document.tables[0]
    clear_body(document)
    append_paragraph(document, source_paragraphs[0], "基本事实")
    append_paragraph(
        document,
        source_paragraphs[1],
        "[用简洁时间线说明与检索问题有关的案件事实；区分已确认事实、当事人陈述、冲突事实和待核实事项。]",
    )
    append_paragraph(document, source_paragraphs[0], "检索问题")
    append_paragraph(document, source_paragraphs[1], "1. [填写需要回答的具体法律问题。]")
    append_paragraph(document, source_paragraphs[1], "2. [填写备选路径、例外或反向问题。]")
    append_paragraph(document, source_paragraphs[3], "法律规定")
    append_paragraph(document, source_paragraphs[4], "《[规范全称]》第[条]条")
    append_paragraph(
        document,
        source_paragraphs[5],
        "第[条]条  [逐字录入经核验的完整条文原文，包含全部款、项、目和标点。]",
    )
    append_paragraph(document, source_paragraphs[4], "《[规范全称]》第[条]条")
    append_paragraph(
        document,
        source_paragraphs[5],
        "第[条]条  [逐字录入经核验的完整条文原文，包含全部款、项、目和标点。]",
    )
    append_paragraph(document, source_paragraphs[50], "案例")
    append_paragraph(document, source_paragraphs[51], "（见下页案例集）")
    append_paragraph(document, source_paragraphs[52], "分析", page_break_before=True)
    append_paragraph(
        document,
        source_paragraphs[53],
        "针对检索问题一：[结合已核验法条、支持与限制性案例以及本案事实，写明分析、条件、风险和初步结论。]",
    )
    append_paragraph(
        document,
        source_paragraphs[53],
        "针对检索问题二：[比较备选方案的成立条件、程序影响、证据缺口和建议。]",
    )
    append_paragraph(document, source_paragraphs[63], "案例一览表", page_break_before=True)
    append_case_table(
        document,
        source_table,
        ["编号", "法院-案号", "裁判要旨", "页码"],
        [
            ["1", "[法院]-（[年份]）[案号]", "[基本案情、争议焦点、裁判理由、裁判结论及与检索问题的对应关系。]", "[页码]"],
            ["2", "[法院]-（[年份]）[案号]", "[限制、相反观点或事实差异，以及本案可借鉴程度。]", "[页码]"],
        ],
    )
    save_sanitized(document, output, "法律检索报告四列案例表脱敏母版")


def build_six_column(source: Path, output: Path) -> None:
    document = Document(source)
    source_paragraphs = list(document.paragraphs)
    source_table = document.tables[0]
    clear_body(document)
    append_paragraph(document, source_paragraphs[0], "案情简要")
    append_paragraph(
        document,
        source_paragraphs[1],
        "[用简洁时间线说明与检索问题有关的案件事实；区分已确认事实、当事人陈述、冲突事实和待核实事项。]",
    )
    append_paragraph(document, source_paragraphs[0], "检索问题")
    append_paragraph(document, source_paragraphs[1], "1. [填写需要回答的具体法律问题。]")
    append_paragraph(document, source_paragraphs[3], "法律规定")
    append_paragraph(document, source_paragraphs[4], "[制定机关]  [规范名称]")
    append_paragraph(
        document,
        source_paragraphs[5],
        "第[条]条  [逐字录入经核验的完整条文原文，包含全部款、项、目和标点。]",
    )
    append_paragraph(document, source_paragraphs[23], "案例")
    append_paragraph(document, source_paragraphs[24], "（见下页案例集）")
    append_paragraph(document, source_paragraphs[25], "分析", page_break_before=True)
    append_paragraph(
        document,
        source_paragraphs[26],
        "[围绕检索问题，将法条、案例规则、事实分支、举证责任、风险和初步建议逐项对应。]",
    )
    append_paragraph(document, source_paragraphs[28], "案例一览表", page_break_before=True)
    append_case_table(
        document,
        source_table,
        ["编号", "法院-案号", "裁判日期", "裁判要旨", "裁判结果", "页码"],
        [
            ["1", "[法院]-（[年份]）[案号]", "[日期]", "[完整记录基本案情、争议焦点、裁判理由及与检索问题的对应关系。]", "[支持/不支持/发回等]", "[页码]"],
            ["2", "[法院]-（[年份]）[案号]", "[日期]", "[记录限制、相反观点或事实差异及可借鉴程度。]", "[裁判结果]", "[页码]"],
        ],
    )
    save_sanitized(document, output, "法律检索报告六列案例表脱敏母版")


def build_law_compilation(source: Path, output: Path) -> None:
    document = Document(source)
    source_paragraphs = list(document.paragraphs)
    clear_body(document)
    append_paragraph(document, source_paragraphs[0], "法律规定整理")
    append_multirun(
        document,
        source_paragraphs[1],
        ["[争点或权利基础一]的法律依据", "【", "原则性的法律依据", "】"],
    )
    append_paragraph(document, source_paragraphs[2], "《[规范全称]》第[条]条")
    append_paragraph(
        document,
        source_paragraphs[3],
        "第[条]条  [逐字录入经核验的完整条文原文，包含全部款、项、目和标点。]",
    )
    append_paragraph(
        document,
        source_paragraphs[4],
        "[如该条包含多款、项、目，继续逐段完整列示，不摘录、不省略。]",
    )
    append_paragraph(document, source_paragraphs[8], "[争点或诉讼请求二]的法律依据")
    append_multirun(
        document,
        source_paragraphs[9],
        ["【对应", "诉请第[ ]项", "“", "[填写诉请内容]", "”】"],
    )
    append_paragraph(document, source_paragraphs[10], "《[规范全称]》第[条]条")
    append_paragraph(
        document,
        source_paragraphs[11],
        "第[条]条  [逐字录入经核验的完整条文原文，包含全部款、项、目和标点。]",
    )
    append_paragraph(document, source_paragraphs[20], "[争点或诉讼请求三]的法律依据")
    append_multirun(
        document,
        source_paragraphs[21],
        ["【", "对应诉请第[ ]项", "“", "[填写诉请内容]", "”】"],
    )
    append_paragraph(document, source_paragraphs[10], "《[规范全称]》第[条]条")
    append_paragraph(
        document,
        source_paragraphs[11],
        "第[条]条  [逐字录入经核验的完整条文原文，包含全部款、项、目和标点。]",
    )
    save_sanitized(document, output, "法律规定整理诉请争点对应脱敏母版")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--four-column-source", type=Path, required=True)
    parser.add_argument("--six-column-source", type=Path, required=True)
    parser.add_argument("--law-compilation-source", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()

    build_four_column(
        args.four_column_source,
        args.output_root / "legal_research" / "TPL-RES-001_法律检索报告_四列案例表_脱敏母版.docx",
    )
    build_six_column(
        args.six_column_source,
        args.output_root / "legal_research" / "TPL-RES-002_法律检索报告_六列案例表_脱敏母版.docx",
    )
    build_law_compilation(
        args.law_compilation_source,
        args.output_root / "legal_authority" / "TPL-LAW-003_法律规定整理_诉请争点对应_脱敏母版.docx",
    )


if __name__ == "__main__":
    main()
